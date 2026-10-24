
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.chart import LineChart, Reference
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
PROCESSED_DIR = PROJECT_ROOT / "data" / "processed"
EXCEL_DIR = PROJECT_ROOT / "excel"
OUTPUT_FILE = EXCEL_DIR / "benchmark_drift_analysis.xlsx"


def style_sheet(ws):
    """Apply readable formatting to a worksheet."""
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    for column_cells in ws.columns:
        max_length = max(
            len(str(cell.value)) if cell.value is not None else 0
            for cell in column_cells
        )
        ws.column_dimensions[
            get_column_letter(column_cells[0].column)
        ].width = min(max(max_length + 2, 12), 34)


def write_dataframe(ws, df):
    ws.append(list(df.columns))

    for row in df.itertuples(index=False, name=None):
        ws.append([
            value.to_pydatetime()
            if isinstance(value, pd.Timestamp)
            else value
            for value in row
        ])

    style_sheet(ws)


def build_workbook():
    aligned_path = PROCESSED_DIR / "aligned_daily_data.csv"
    metrics_path = PROCESSED_DIR / "summary_metrics.csv"
    episodes_path = PROCESSED_DIR / "stress_episodes.csv"

    for path in (aligned_path, metrics_path, episodes_path):
        if not path.exists():
            raise FileNotFoundError(
                f"{path} not found. Run `python main.py` first."
            )

    data = pd.read_csv(aligned_path, parse_dates=["date"])
    metrics_df = pd.read_csv(metrics_path)
    episodes_df = pd.read_csv(
        episodes_path,
        parse_dates=["start_date", "end_date"],
    )

    EXCEL_DIR.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    dashboard = wb.active
    dashboard.title = "Dashboard"
    dashboard.append(["Passive Fund Tracking Stress & Benchmark Drift Analysis"])
    dashboard["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    dashboard["A1"].fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )
    dashboard.merge_cells("A1:B1")

    dashboard.append(["Metric", "Value"])

    for row in metrics_df.itertuples(index=False, name=None):
        dashboard.append([row[0], row[1]])

    style_sheet(dashboard)

    aligned_ws = wb.create_sheet("Aligned Data")
    write_dataframe(aligned_ws, data)

    # Format dates and returns.
    for row in range(2, aligned_ws.max_row + 1):
        aligned_ws.cell(row, 1).number_format = "yyyy-mm-dd"

    for column_name in (
        "fund_return",
        "benchmark_return",
        "active_return",
        "rolling_60d_tracking_error",
    ):
        if column_name in data.columns:
            col = data.columns.get_loc(column_name) + 1
            for row in range(2, aligned_ws.max_row + 1):
                aligned_ws.cell(row, col).number_format = "0.000000%"

    stress_ws = wb.create_sheet("Stress Episodes")
    write_dataframe(stress_ws, episodes_df)

    for row in range(2, stress_ws.max_row + 1):
        stress_ws.cell(row, 2).number_format = "yyyy-mm-dd"
        stress_ws.cell(row, 3).number_format = "yyyy-mm-dd"

    methodology = wb.create_sheet("Methodology")
    methodology.append(["Item", "Definition / Note"])
    methodology_rows = [
        (
            "Fund",
            "Kotak Nifty 50 Index Fund - Direct Plan - Growth; "
            "AMFI scheme code 148978.",
        ),
        (
            "Benchmark",
            "NIFTY 50 Total Returns Index (TRI); verify official source "
            "and selected series before final publication.",
        ),
        (
            "Daily fund return",
            "Current NAV divided by previous aligned NAV, minus 1.",
        ),
        (
            "Daily benchmark return",
            "Current TRI divided by previous aligned TRI, minus 1.",
        ),
        (
            "Active return",
            "Daily fund return minus daily benchmark return.",
        ),
        (
            "Tracking error",
            "Sample standard deviation of daily active returns "
            "multiplied by square root of 252.",
        ),
        (
            "Information ratio",
            "Mean daily active return multiplied by 252, divided "
            "by annualized tracking error.",
        ),
        (
            "Rolling tracking error",
            "Sample standard deviation over 60 aligned return "
            "observations, annualized by square root of 252.",
        ),
        (
            "Stress threshold",
            "95th percentile of observed rolling tracking error; "
            "retrospective descriptive threshold, not a universal standard.",
        ),
        (
            "Limitations",
            "Observed deviations do not establish causation. "
            "No inference about expenses, cash holdings, or trading costs "
            "is made without supporting evidence.",
        ),
    ]

    for row in methodology_rows:
        methodology.append(list(row))

    style_sheet(methodology)

    # Add a cumulative growth chart to the dashboard.
    if {"fund_return", "benchmark_return"} <= set(data.columns):
        chart_data = data.copy()
        chart_data["fund_growth"] = (
            1 + chart_data["fund_return"].fillna(0)
        ).cumprod()
        chart_data["benchmark_growth"] = (
            1 + chart_data["benchmark_return"].fillna(0)
        ).cumprod()

        aligned_ws.cell(1, aligned_ws.max_column + 1, "fund_growth")
        aligned_ws.cell(1, aligned_ws.max_column + 1, "benchmark_growth")

        fund_col = aligned_ws.max_column - 1
        benchmark_col = aligned_ws.max_column

        for idx, row in enumerate(
            chart_data.itertuples(index=False), start=2
        ):
            aligned_ws.cell(idx, fund_col, row[-2])
            aligned_ws.cell(idx, benchmark_col, row[-1])

        chart = LineChart()
        chart.title = "Cumulative Growth"
        chart.y_axis.title = "Growth of initial investment"
        chart.x_axis.title = "Date"
        chart.height = 8
        chart.width = 16

        values = Reference(
            aligned_ws,
            min_col=fund_col,
            max_col=benchmark_col,
            min_row=1,
            max_row=aligned_ws.max_row,
        )
        dates = Reference(
            aligned_ws,
            min_col=1,
            min_row=2,
            max_row=aligned_ws.max_row,
        )
        chart.add_data(values, titles_from_data=True)
        chart.set_categories(dates)
        dashboard.add_chart(chart, "D2")

    wb.save(OUTPUT_FILE)
    print(f"Excel workbook created: {OUTPUT_FILE}")


if __name__ == "__main__":
    build_workbook()